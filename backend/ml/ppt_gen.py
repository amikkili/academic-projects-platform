from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from .slide_content import SLIDE_CONTENT

# ── Brand colours ─────────────────────────────────────────────────────────────
NAVY   = RGBColor(0x0B, 0x1D, 0x3A)
TEAL   = RGBColor(0x1A, 0xBA, 0xAF)
ORANGE = RGBColor(0xF9, 0x73, 0x16)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x1E, 0x29, 0x3B)
GRAY   = RGBColor(0x64, 0x74, 0x8B)
LIGHT  = RGBColor(0xF1, 0xF5, 0xF9)

# ── Layout constants (widescreen 13.33 × 7.5 in) ─────────────────────────────
W = Inches(13.33)
H = Inches(7.5)
BAR_H    = Inches(1.25)
FOOT_H   = Inches(0.38)
FOOT_TOP = H - FOOT_H
BODY_TOP = BAR_H + Inches(0.15)
BODY_H   = FOOT_TOP - BODY_TOP - Inches(0.1)
PAD      = Inches(0.5)


# ── Helper: solid rectangle ───────────────────────────────────────────────────
def _rect(slide, left, top, width, height, fill_rgb, border=False):
    shape = slide.shapes.add_shape(1, left, top, width, height)  # 1 = rectangle
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    if border:
        shape.line.color.rgb = fill_rgb
    else:
        shape.line.fill.background()
    return shape


# ── Helper: text box ──────────────────────────────────────────────────────────
def _textbox(slide, text, left, top, width, height,
             size=16, bold=False, color=DARK, align=PP_ALIGN.LEFT, italic=False):
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf  = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size    = Pt(size)
    run.font.bold    = bold
    run.font.italic  = italic
    run.font.color.rgb = color
    return txb


# ── Helper: content slide header bar ─────────────────────────────────────────
def _header(slide, title, slide_num, total, project_name):
    _rect(slide, 0, 0, W, BAR_H, NAVY)
    # left teal accent strip
    _rect(slide, 0, 0, Inches(0.07), BAR_H, TEAL)
    # slide title
    _textbox(slide, title,
             Inches(0.25), Inches(0.2), Inches(11.5), Inches(0.85),
             size=24, bold=True, color=WHITE)
    # slide counter top-right
    _textbox(slide, f"{slide_num} / {total}",
             Inches(11.8), Inches(0.45), Inches(1.4), Inches(0.4),
             size=11, color=RGBColor(0xAA, 0xBB, 0xCC), align=PP_ALIGN.RIGHT)


# ── Helper: footer bar ────────────────────────────────────────────────────────
def _footer(slide, project_name):
    _rect(slide, 0, FOOT_TOP, W, FOOT_H, LIGHT)
    _rect(slide, 0, FOOT_TOP, W, Inches(0.03), TEAL)
    _textbox(slide, project_name,
             PAD, FOOT_TOP + Inches(0.07), Inches(9), Inches(0.28),
             size=9, color=GRAY)
    _textbox(slide, "Academic Project Presentation",
             Inches(9.5), FOOT_TOP + Inches(0.07), Inches(3.7), Inches(0.28),
             size=9, color=GRAY, align=PP_ALIGN.RIGHT)


# ── Helper: bullet list in body area ─────────────────────────────────────────
def _bullets(slide, points, top_offset=0):
    txb = slide.shapes.add_textbox(
        PAD,
        BODY_TOP + Inches(top_offset),
        W - PAD * 2,
        BODY_H - Inches(top_offset),
    )
    tf = txb.text_frame
    tf.word_wrap = True

    for i, point in enumerate(points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(point, tuple):
            # (label, description) → bold label + normal desc
            run1 = p.add_run()
            run1.text = f"▸  {point[0]}: "
            run1.font.size  = Pt(15)
            run1.font.bold  = True
            run1.font.color.rgb = NAVY
            run2 = p.add_run()
            run2.text = point[1]
            run2.font.size  = Pt(15)
            run2.font.bold  = False
            run2.font.color.rgb = DARK
        else:
            run = p.add_run()
            run.text = f"▸  {point}"
            run.font.size   = Pt(15)
            run.font.bold   = False
            run.font.color.rgb = DARK


# ── Slide builders ────────────────────────────────────────────────────────────

def _slide_title(prs, content, student_name, branch, year):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # Full navy background
    _rect(slide, 0, 0, W, H, NAVY)
    # Top teal accent bar
    _rect(slide, 0, 0, W, Inches(0.12), TEAL)
    # Bottom teal accent bar
    _rect(slide, 0, H - Inches(0.12), W, Inches(0.12), TEAL)
    # Orange decorative stripe
    _rect(slide, 0, Inches(0.12), Inches(0.08), H - Inches(0.24), ORANGE)

    # Project title
    txb = slide.shapes.add_textbox(Inches(1.0), Inches(1.4), Inches(11.3), Inches(1.8))
    tf  = txb.text_frame
    tf.word_wrap = True
    p   = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = content['title']
    run.font.size  = Pt(36)
    run.font.bold  = True
    run.font.color.rgb = WHITE

    # Subtitle
    _textbox(slide, content['subtitle'],
             Inches(1.0), Inches(3.3), Inches(11.3), Inches(0.7),
             size=18, color=TEAL, align=PP_ALIGN.CENTER)

    # Domain tag
    _textbox(slide, f"Domain: {content['domain']}",
             Inches(1.0), Inches(4.1), Inches(11.3), Inches(0.5),
             size=14, color=RGBColor(0xCC, 0xDD, 0xEE), align=PP_ALIGN.CENTER)

    # Divider line
    line = slide.shapes.add_shape(1, Inches(3.5), Inches(4.7), Inches(6.3), Inches(0.04))
    line.fill.solid(); line.fill.fore_color.rgb = TEAL; line.line.fill.background()

    # Student info
    info_lines = []
    if student_name: info_lines.append(f"Presented by: {student_name}")
    if branch:       info_lines.append(f"Branch: {branch}")
    if year:         info_lines.append(f"Academic Year: {year}")

    for j, line_text in enumerate(info_lines):
        _textbox(slide, line_text,
                 Inches(1.0), Inches(5.0 + j * 0.45), Inches(11.3), Inches(0.42),
                 size=14, color=RGBColor(0xCC, 0xDD, 0xEE), align=PP_ALIGN.CENTER)


def _slide_agenda(prs, content, slide_num, total, project_name):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "Agenda", slide_num, total, project_name)
    _footer(slide, project_name)

    txb = slide.shapes.add_textbox(PAD, BODY_TOP, W - PAD * 2, BODY_H)
    tf  = txb.text_frame
    tf.word_wrap = True

    # Two-column layout using two text boxes
    items = content['agenda']
    mid = (len(items) + 1) // 2
    left_items  = items[:mid]
    right_items = items[mid:]

    for col_items, left_pos in [(left_items, PAD), (right_items, Inches(7.0))]:
        col = slide.shapes.add_textbox(left_pos, BODY_TOP + Inches(0.1), Inches(5.8), BODY_H)
        ctf = col.text_frame
        ctf.word_wrap = True
        for i, item in enumerate(col_items):
            p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
            run = p.add_run()
            num = items.index(item) + 1
            run.text = f"{num:02d}.  {item}"
            run.font.size  = Pt(15)
            run.font.bold  = False
            run.font.color.rgb = DARK


def _slide_content(prs, title, points, slide_num, total, project_name):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, title, slide_num, total, project_name)
    _footer(slide, project_name)
    _bullets(slide, points)


def _slide_tech_stack(prs, content, slide_num, total, project_name):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, content['tech_stack']['title'], slide_num, total, project_name)
    _footer(slide, project_name)

    items = content['tech_stack']['items']
    mid   = (len(items) + 1) // 2
    left_items  = items[:mid]
    right_items = items[mid:]

    for col_items, left_pos in [(left_items, PAD), (right_items, Inches(7.0))]:
        col = slide.shapes.add_textbox(left_pos, BODY_TOP + Inches(0.15), Inches(5.8), BODY_H)
        ctf = col.text_frame
        ctf.word_wrap = True
        for i, (tech, desc) in enumerate(col_items):
            p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
            r1 = p.add_run()
            r1.text = f"▸  {tech}"
            r1.font.size  = Pt(15)
            r1.font.bold  = True
            r1.font.color.rgb = NAVY
            # description on next line via a sub-paragraph
            p2 = ctf.add_paragraph()
            r2 = p2.add_run()
            r2.text = f"     {desc}"
            r2.font.size   = Pt(13)
            r2.font.italic = True
            r2.font.color.rgb = GRAY


def _slide_two_col(prs, title, left_title, left_pts, right_title, right_pts,
                   slide_num, total, project_name, left_color=TEAL, right_color=ORANGE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, title, slide_num, total, project_name)
    _footer(slide, project_name)

    # Left column header
    _rect(slide, PAD, BODY_TOP, Inches(5.8), Inches(0.45), left_color)
    _textbox(slide, left_title,
             PAD + Inches(0.1), BODY_TOP + Inches(0.07), Inches(5.6), Inches(0.35),
             size=14, bold=True, color=WHITE)

    # Right column header
    _rect(slide, Inches(7.0), BODY_TOP, Inches(5.8), Inches(0.45), right_color)
    _textbox(slide, right_title,
             Inches(7.1), BODY_TOP + Inches(0.07), Inches(5.6), Inches(0.35),
             size=14, bold=True, color=WHITE)

    # Left bullets
    left_top = BODY_TOP + Inches(0.55)
    left_txb = slide.shapes.add_textbox(PAD, left_top, Inches(5.8), BODY_H - Inches(0.55))
    ltf = left_txb.text_frame; ltf.word_wrap = True
    for i, pt in enumerate(left_pts):
        p = ltf.paragraphs[0] if i == 0 else ltf.add_paragraph()
        run = p.add_run(); run.text = f"✔  {pt}"
        run.font.size = Pt(14); run.font.color.rgb = DARK

    # Right bullets
    right_txb = slide.shapes.add_textbox(Inches(7.0), left_top, Inches(5.8), BODY_H - Inches(0.55))
    rtf = right_txb.text_frame; rtf.word_wrap = True
    for i, pt in enumerate(right_pts):
        p = rtf.paragraphs[0] if i == 0 else rtf.add_paragraph()
        run = p.add_run(); run.text = f"✗  {pt}"
        run.font.size = Pt(14); run.font.color.rgb = DARK


def _slide_references(prs, refs, slide_num, total, project_name):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "References", slide_num, total, project_name)
    _footer(slide, project_name)
    txb = slide.shapes.add_textbox(PAD, BODY_TOP + Inches(0.1), W - PAD * 2, BODY_H)
    tf  = txb.text_frame; tf.word_wrap = True
    for i, ref in enumerate(refs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r1 = p.add_run(); r1.text = f"[{i+1}]  "
        r1.font.size = Pt(13); r1.font.bold = True; r1.font.color.rgb = NAVY
        r2 = p.add_run(); r2.text = ref
        r2.font.size = Pt(13); r2.font.color.rgb = DARK


def _slide_thankyou(prs, content, project_name):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, W, H, NAVY)
    _rect(slide, 0, 0, W, Inches(0.12), TEAL)
    _rect(slide, 0, H - Inches(0.12), W, Inches(0.12), TEAL)
    _rect(slide, 0, Inches(0.12), Inches(0.08), H - Inches(0.24), ORANGE)

    _textbox(slide, "Thank You",
             Inches(1.0), Inches(1.8), Inches(11.3), Inches(1.5),
             size=48, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    _textbox(slide, content['title'],
             Inches(1.0), Inches(3.5), Inches(11.3), Inches(0.7),
             size=20, color=TEAL, align=PP_ALIGN.CENTER)

    _textbox(slide, "We welcome your questions and feedback",
             Inches(1.0), Inches(4.35), Inches(11.3), Inches(0.55),
             size=15, color=RGBColor(0xCC, 0xDD, 0xEE), align=PP_ALIGN.CENTER,
             italic=True)

    line = slide.shapes.add_shape(1, Inches(3.5), Inches(5.0), Inches(6.3), Inches(0.04))
    line.fill.solid(); line.fill.fore_color.rgb = TEAL; line.line.fill.background()

    _textbox(slide, "AcademiCode | Academic Project Experts",
             Inches(1.0), Inches(5.2), Inches(11.3), Inches(0.45),
             size=13, color=GRAY, align=PP_ALIGN.CENTER)


# ── Public entry point ────────────────────────────────────────────────────────

def build_ppt(project_id: str, student_name: str = "", branch: str = "", year: str = "") -> BytesIO:
    """
    Build a .pptx for the given project and return it as a BytesIO stream.

    Args:
        project_id:   slug from SLIDE_CONTENT (e.g. 'sentiment-analysis')
        student_name: shown on the title slide (optional)
        branch:       e.g. 'B.Tech – CSE' (optional)
        year:         e.g. '2024–25' (optional)

    Returns:
        BytesIO containing the .pptx binary — seek to 0 before sending.
    Raises:
        ValueError if project_id is not in SLIDE_CONTENT.
    """
    content = SLIDE_CONTENT.get(project_id)
    if content is None:
        raise ValueError(f"No slide content for project: {project_id!r}")

    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H

    # Count total slides: title + agenda + 11 content + thank you = 14
    TOTAL = 14

    # 1. Title slide
    _slide_title(prs, content, student_name, branch, year)

    # 2. Agenda
    _slide_agenda(prs, content, 2, TOTAL, content['title'])

    # 3. Introduction
    _slide_content(prs, content['introduction']['title'],
                   content['introduction']['points'], 3, TOTAL, content['title'])

    # 4. Objectives
    _slide_content(prs, content['objectives']['title'],
                   content['objectives']['points'], 4, TOTAL, content['title'])

    # 5. Architecture
    _slide_content(prs, content['architecture']['title'],
                   content['architecture']['points'], 5, TOTAL, content['title'])

    # 6. Tech Stack (two-column layout)
    _slide_tech_stack(prs, content, 6, TOTAL, content['title'])

    # 7. Module 1 Implementation
    _slide_content(prs, content['module1']['title'],
                   content['module1']['points'], 7, TOTAL, content['title'])

    # 8. Module 2 Implementation
    _slide_content(prs, content['module2']['title'],
                   content['module2']['points'], 8, TOTAL, content['title'])

    # 9. Results & Output
    _slide_content(prs, content['results']['title'],
                   content['results']['points'], 9, TOTAL, content['title'])

    # 10. Advantages & Limitations (two-column)
    _slide_two_col(prs,
                   "Advantages & Limitations",
                   "Advantages", content['advantages'],
                   "Limitations", content['limitations'],
                   10, TOTAL, content['title'],
                   left_color=TEAL, right_color=ORANGE)

    # 11. Future Scope
    _slide_content(prs, "Future Scope", content['future_scope'], 11, TOTAL, content['title'])

    # 12. Conclusion
    _slide_content(prs, "Conclusion", content['conclusion'], 12, TOTAL, content['title'])

    # 13. References
    _slide_references(prs, content['references'], 13, TOTAL, content['title'])

    # 14. Thank You
    _slide_thankyou(prs, content, content['title'])

    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf
